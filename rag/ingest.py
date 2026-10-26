# Add every supported file in rag/documents/ to the vector store (knowledge.db).
#
# Usage:
#   uv run rag/ingest.py             add new files, re-embed changed ones, remove deleted ones
#   uv run rag/ingest.py --rebuild   delete knowledge.db and embed everything again
#
# Supported: .pdf .docx .pptx .txt .md .csv .html .htm
# Files can be on any topic and can sit in subfolders; folder names don't matter.

import sys, csv, hashlib, argparse
from datetime import datetime

from rag_common import (DOCS_DIR, DB_PATH, CHUNK_WORDS, OVERLAP_WORDS,
                        embed_texts, open_db, to_blob)


# ---------- Loaders: each returns (kind, [(label, text), ...]) ----------
# "kind" names the location unit (page, slide, line, ...); "label" is that
# unit's number, or a section heading for Word files.

def load_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "page", [(i, page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]


def load_docx(path):
    from docx import Document
    doc = Document(path)
    paras = [p for p in doc.paragraphs if p.text.strip()]

    def is_heading(p):
        name = (p.style.name if p.style is not None else "").lower()
        return name.startswith("heading") or name == "title"

    units = []
    if any(is_heading(p) for p in paras):
        # Label each paragraph with the heading it falls under
        kind, section = "section", "Beginning"
        for p in paras:
            if is_heading(p):
                section = p.text.strip()[:60]
            units.append((section, p.text))
    else:
        # No headings: number the paragraphs instead
        kind = "paragraph"
        units = [(i, p.text) for i, p in enumerate(paras, 1)]

    # Tables (text only, one row per line)
    for t, table in enumerate(doc.tables, 1):
        label = f"Table {t}" if kind == "section" else len(units) + 1
        for row in table.rows:
            cells = []
            for c in row.cells:                      # merged cells repeat, skip duplicates
                txt = c.text.strip()
                if txt and (not cells or cells[-1] != txt):
                    cells.append(txt)
            if cells:
                units.append((label, " | ".join(cells)))
    return kind, units


def load_pptx(path):
    from pptx import Presentation
    units = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        units.append((i, "\n".join(parts)))
    return "slide", units


def read_text(path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def load_txt(path):
    return "line", list(enumerate(read_text(path).splitlines(), 1))


def load_csv(path):
    rows = list(csv.reader(read_text(path).splitlines()))
    if not rows:
        return "row", []
    header, units = rows[0], []
    for i, row in enumerate(rows[1:], 1):
        # "column: value" keeps the meaning of each cell for the search
        pairs = [f"{h}: {v}" for h, v in zip(header, row) if v.strip()]
        units.append((i, "; ".join(pairs) + "."))
    return "row", units


def load_html(path):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(read_text(path), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    lines = [l.strip() for l in soup.get_text("\n").splitlines() if l.strip()]
    return "line", list(enumerate(lines, 1))


LOADERS = {
    ".pdf": load_pdf, ".docx": load_docx, ".pptx": load_pptx,
    ".txt": load_txt, ".md": load_txt, ".csv": load_csv,
    ".html": load_html, ".htm": load_html,
}


# ---------- Chunking ----------
def format_location(kind, labels):
    labels = list(dict.fromkeys(labels))            # unique, in order
    if all(isinstance(x, int) for x in labels):
        lo, hi = min(labels), max(labels)
        return f"{kind} {lo}" if lo == hi else f"{kind}s {lo}-{hi}"
    shown = "; ".join(str(x) for x in labels[:3])
    return f"{kind}: {shown}" + (" ..." if len(labels) > 3 else "")


def make_chunks(kind, units):
    """Sliding window of CHUNK_WORDS words with OVERLAP_WORDS overlap.
    Each word remembers its label, so each chunk knows where it came from."""
    words, labels = [], []
    for label, text in units:
        w = text.split()
        words += w
        labels += [label] * len(w)

    chunks, start, step = [], 0, CHUNK_WORDS - OVERLAP_WORDS
    while start < len(words):
        end = min(start + CHUNK_WORDS, len(words))
        chunks.append((format_location(kind, labels[start:end]),
                       " ".join(words[start:end])))
        if end == len(words):
            break
        start += step
    return chunks, len(words)


# ---------- Store maintenance ----------
def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def remove_file(db, file_id):
    ids = [r[0] for r in db.execute("SELECT id FROM chunks WHERE file_id = ?", (file_id,))]
    for cid in ids:
        db.execute("DELETE FROM vec_chunks WHERE rowid = ?", (cid,))
    db.execute("DELETE FROM chunks WHERE file_id = ?", (file_id,))
    db.execute("DELETE FROM files WHERE id = ?", (file_id,))


def add_file(db, path, rel, digest):
    kind, units = LOADERS[path.suffix.lower()](path)
    chunks, n_words = make_chunks(kind, units)
    if n_words < 20:
        print(f"  skipped: almost no text found "
              f"({'scanned PDF? needs OCR' if path.suffix.lower() == '.pdf' else 'empty file?'})")
        db.commit()
        return 0

    print(f"  {n_words} words -> {len(chunks)} chunks, embedding...")
    vectors = embed_texts([text for _, text in chunks])   # before writing anything

    # Write the whole file in one transaction: a crash never leaves half a file.
    cur = db.execute(
        "INSERT INTO files (path, hash, file_type, chunks, added) VALUES (?, ?, ?, ?, ?)",
        (rel, digest, path.suffix.lower().lstrip("."), len(chunks),
         datetime.now().isoformat(timespec="seconds")))
    file_id = cur.lastrowid
    for (location, text), vec in zip(chunks, vectors):
        cid = db.execute("INSERT INTO chunks (file_id, location, text) VALUES (?, ?, ?)",
                         (file_id, location, text)).lastrowid
        db.execute("INSERT INTO vec_chunks (rowid, embedding) VALUES (?, ?)", (cid, to_blob(vec)))
    db.commit()
    return len(chunks)


def main():
    parser = argparse.ArgumentParser(description="Add documents to knowledge.db")
    parser.add_argument("--rebuild", action="store_true", help="start from an empty store")
    args = parser.parse_args()

    if args.rebuild and DB_PATH.exists():
        DB_PATH.unlink()
        print("Deleted old knowledge.db")

    DOCS_DIR.mkdir(exist_ok=True)
    db = open_db()

    on_disk = {}
    for p in sorted(DOCS_DIR.rglob("*")):
        if not p.is_file() or p.name.startswith(("~$", ".")):
            continue                                       # skip Office temp files
        if p.suffix.lower() in LOADERS:
            on_disk[p.relative_to(DOCS_DIR).as_posix()] = p
        else:
            print(f"Unsupported file type, ignored: {p.name}")

    known = {path: (fid, h) for fid, path, h in db.execute("SELECT id, path, hash FROM files")}

    # Files that were deleted from the folder
    for rel, (fid, _) in known.items():
        if rel not in on_disk:
            remove_file(db, fid)
            db.commit()
            print(f"Removed (no longer in documents): {rel}")

    added = updated = unchanged = failed = 0
    for rel, path in on_disk.items():
        digest = file_hash(path)
        if rel in known and known[rel][1] == digest:
            unchanged += 1
            continue
        print(f"{'Updating' if rel in known else 'Adding'}: {rel}")
        try:
            if rel in known:
                remove_file(db, known[rel][0])
            if add_file(db, path, rel, digest):
                updated += rel in known
                added += rel not in known
        except Exception as e:
            db.rollback()
            failed += 1
            print(f"  failed to read this file: {e}")

    total_files, total_chunks = db.execute(
        "SELECT COUNT(*), COALESCE(SUM(chunks), 0) FROM files").fetchone()
    print(f"\nDone. Added {added}, updated {updated}, unchanged {unchanged}, failed {failed}.")
    print(f"The store now holds {total_files} files and {total_chunks} chunks.")
    if not on_disk:
        print(f"No documents found. Put your files in: {DOCS_DIR}")


if __name__ == "__main__":
    main()